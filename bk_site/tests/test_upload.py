import os
import openpyxl
from bk_site.data_loader import load_dataset
from bk_site.dataset_store import validate_xlsx

SRC = "E:/AI Agent/Riset_22Sep/bk_site/uploads/Data Kelas 5 lengkap.xlsx"
TMP = "C:/Users/pandu/AppData/Local/Temp/opencode"


def _variant(path, sheet_names=None, shuffle=False, drop_col=None):
    w = openpyxl.load_workbook(SRC)
    names = sheet_names or list(w.sheetnames)
    for old, new in zip(list(w.sheetnames), names):
        w[old].title = new
    if drop_col:
        for ws in w.worksheets:
            head = [c.value for c in ws[1]]
            if drop_col in head:
                ws.delete_cols(head.index(drop_col) + 1)
                break
    if shuffle:
        ws = w[names[0]]
        rows = list(ws.iter_rows(values_only=True))
        head, data = rows[0], rows[1:]
        order = list(range(len(head)))
        order = order[2:] + order[:2]
        ws.delete_rows(1, ws.max_row)
        ws.append([head[i] for i in order])
        for r in data:
            ws.append([r[i] for i in order])
    w.save(path)
    return path


def test_sheet_bebas_dan_kolom_acak():
    p = os.path.join(TMP, "varian.xlsx")
    _variant(p, sheet_names=["DataNilai", "Rekap"], shuffle=True)
    assert validate_xlsx(p) == []
    d = load_dataset(p)
    assert len(d["students"]) == 24
    s = [x for x in d["students"] if x["nama"] == "ALYA ZHAFIRAH"][0]
    assert s["kelas_rows"]["Kelas 5"]["rata"] == 94
    assert s["tren"] == "Naik"


def test_kolom_kurang_ditolak_ramah():
    p = os.path.join(TMP, "kurang.xlsx")
    _variant(p, drop_col="poin_pelanggaran")
    errs = validate_xlsx(p)
    assert errs and any("nilai" in e for e in errs)


def test_tanpa_kolom_tren_diterima():
    p = os.path.join(TMP, "notren.xlsx")
    _variant(p, drop_col="Tren Akhir")
    assert validate_xlsx(p) == []
    d = load_dataset(p)
    assert d["students"][0]["tren"] in ("Naik", "Turun")


def test_hapus_data():
    import shutil
    up = "E:/AI Agent/Riset_22Sep/bk_site/uploads"
    shutil.copy(os.path.join(up, "Data Kelas 5 lengkap.xlsx"), os.path.join(up, "hapus_saya.xlsx"))
    from bk_site.app import app
    c = app.test_client()
    assert c.post("/data/hapus", data={"file": "hapus_saya.xlsx"}).status_code == 302
    assert not os.path.exists(os.path.join(up, "hapus_saya.xlsx"))
    r = c.post("/data/hapus", data={"file": "Data Kelas 5 lengkap.xlsx"})
    assert r.status_code == 200 and "tak boleh dihapus" in r.data.decode()
