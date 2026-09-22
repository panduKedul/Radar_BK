import os
import openpyxl
from bk_site.data_loader import load_dataset
from bk_site.risk import compute_risk

SRC = "E:/AI Agent/Riset_22Sep/bk_site/uploads/Data Kelas 5 lengkap.xlsx"
TMP = "C:/Users/pandu/AppData/Local/Temp/opencode"


def _kelas_b(path):
    w = openpyxl.load_workbook(SRC)
    ws = w["Dataset_Format_Website"]
    for row in ws.iter_rows(min_row=2, values_only=False):
        row[2].value = str(row[2].value).replace("Kelas ", "Kelas ") + " B"
    w.remove(w["Rekap_Per_Mapel_K1_sd_K5"])
    w.create_sheet("Rekap_Per_Mapel_K1_sd_K5").append(
        ["No", "ID Siswa", "Nama Siswa", "IPAS K1", "IPAS K2", "IPAS K3",
         "IPAS K4", "IPAS K5 (Asli)", "B.Ind K1", "B.Ind K2", "B.Ind K3",
         "B.Ind K4", "B.Ind K5 (Asli)", "MTK K1", "MTK K2", "MTK K3",
         "MTK K4", "MTK K5 (Asli)", "Rata-rata K5 (Asli)", "Tren Akhir"])
    w.save(path)
    return path


def test_kelas_label_bebas_dan_rekap_kosong():
    p = os.path.join(TMP, "kelasb.xlsx")
    _kelas_b(p)
    d = load_dataset(p)
    assert len(d["students"]) == 24
    s = d["students"][0]
    assert s["kelas_akhir"] == "Kelas 5 B"
    assert len(s["ip_k"]) == 5 and s["tren"] in ("Naik", "Turun")
    r = compute_risk(s)
    assert r["badge"] in ("Aman", "Pantau", "Intervensi")


def test_k5_asli_terbaca():
    d = load_dataset(SRC)
    s = [x for x in d["students"] if x["nama"] == "ALYA ZHAFIRAH"][0]
    assert s["ip_k"] == [95.48, 97.29, 95.03, 94.39, 95]
    assert s["tren"] == "Naik"


def test_tren_tahan_nilai_kosong():
    import bk_site.app as A
    d = load_dataset(SRC)
    d["students"][0]["ip_k"][4] = None
    A.DATA = d
    try:
        assert A.app.test_client().get("/tren").status_code == 200
    finally:
        A.DATA = load_dataset(SRC)
