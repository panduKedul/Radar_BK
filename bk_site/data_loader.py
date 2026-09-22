"""Muat dataset radar BK dari xlsx 2 sheet (nama sheet & urutan kolom bebas)."""
import openpyxl
import re

REQUIRED1 = ["id", "nama", "kelas", "nilai_ipas", "nilai_bind", "nilai_mtk",
             "absensi harian", "rata_nilai", "poin_pelanggaran"]
REQUIRED2 = ["ID Siswa", "Nama Siswa"]
OPTIONAL2 = ["Rata-rata K5 (Asli)", "Tren Akhir"]


def norm(h):
    return str(h or "").strip().lower()


def kelas_no(k):
    m = re.search(r"(\d+)", k or "")
    return int(m.group(1)) if m else 0


def detect_sheets(w):
    """Kembalikan {sheet_utama, sheet_rekap} berdasar isi header, bukan nama."""
    found = {}
    for sh in w.sheetnames:
        head = [norm(c) for c in next(w[sh].iter_rows(values_only=True))]
        if all(norm(c) in head for c in REQUIRED1):
            found["utama"] = sh
        if all(norm(c) in head for c in REQUIRED2):
            found["rekap"] = sh
    return found


def colmap(w, sheet, required):
    head = [norm(c) for c in next(w[sheet].iter_rows(values_only=True))]
    return {c: head.index(norm(c)) for c in required}


def parse_pct(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).replace("%", "").strip() or 0)


def tren_otomatis(ratas):
    """Naik bila rata akhir >= rata awal (abaikan kosong)."""
    vals = [x for x in ratas if x is not None]
    if len(vals) < 2:
        return "-"
    return "Naik" if vals[-1] >= vals[0] else "Turun"


def load_dataset(path):
    w = openpyxl.load_workbook(path, read_only=True, data_only=True)
    found = detect_sheets(w)
    if "utama" not in found or "rekap" not in found:
        w.close()
        raise ValueError("File butuh 2 bagian: data nilai (id,nama,kelas,3 nilai,absensi,rata,poin) + rekap (ID,Nama,Rata K5,Tren). Unduh template.")
    s = w[found["utama"]]
    m = colmap(w, found["utama"], REQUIRED1)
    rows = list(s.iter_rows(values_only=True))[1:]
    by_id = {}
    for r in rows:
        if not r[m["id"]] or not r[m["nama"]]:
            continue
        sid, nama, kelas = str(r[m["id"]]).strip(), str(r[m["nama"]]).strip(), str(r[m["kelas"]]).strip()
        by_id.setdefault(sid, {"id": sid, "nama": nama, "kelas_rows": {}})
        by_id[sid]["kelas_rows"][kelas] = {
            "ipas": r[m["nilai_ipas"]], "bind": r[m["nilai_bind"]], "mtk": r[m["nilai_mtk"]],
            "absensi": parse_pct(r[m["absensi harian"]]), "rata": r[m["rata_nilai"]],
            "poin": r[m["poin_pelanggaran"]] or 0,
        }
    s2 = w[found["rekap"]]
    head2 = [norm(c) for c in next(s2.iter_rows(values_only=True))]
    m2 = {c: head2.index(norm(c)) for c in REQUIRED2 + OPTIONAL2 if norm(c) in head2}
    h2 = head2
    for r in list(s2.iter_rows(values_only=True))[1:]:
        sid = str(r[m2["ID Siswa"]] or "").strip()
        if sid in by_id:
            row = dict(zip(h2, r))
            def get(*names):
                for n in names:
                    if norm(n) in row:
                        return row[norm(n)]
                return None
            def series(*prefixes):
                out = []
                for i in range(1, 6):
                    cands = []
                    for p in prefixes:
                        if i == 5:
                            cands.append("%s K%d (Asli)" % (p, i))
                        cands.append("%s K%d" % (p, i))
                    out.append(get(*cands))
                return out
            by_id[sid].update({
                "ip_k": series("IPAS"),
                "bi_k": series("B.Ind", "B. Indonesia"),
                "mtk_k": series("MTK", "Matematika"),
                "rata_k5": r[m2["Rata-rata K5 (Asli)"]] if "Rata-rata K5 (Asli)" in m2 else None,
                "tren": str(r[m2["Tren Akhir"]] or "").strip() if "Tren Akhir" in m2 else "",
            })
    w.close()
    for s in by_id.values():
        order = sorted(s["kelas_rows"], key=kelas_no)
        s["kelas_labels"] = order
        s["kelas_akhir"] = order[-1] if order else "-"
        if "ip_k" not in s:
            s["ip_k"] = [s["kelas_rows"][k]["ipas"] for k in order]
            s["bi_k"] = [s["kelas_rows"][k]["bind"] for k in order]
            s["mtk_k"] = [s["kelas_rows"][k]["mtk"] for k in order]
            s["rata_k5"] = (s["kelas_rows"][order[-1]]["rata"] if order else None)
            s["tren"] = ""
        if not s.get("tren"):
            s["tren"] = tren_otomatis([s["kelas_rows"][k]["rata"] for k in order])
    return {"students": list(by_id.values()), "by_id": by_id}
